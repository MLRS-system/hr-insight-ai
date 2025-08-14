# utils/file_handler.py

import os
import pandas as pd
import PyPDF2
import requests
from bs4 import BeautifulSoup
from youtube_transcript_api import YouTubeTranscriptApi
import docx
from pptx import Presentation

# --- Web Content Extractors ---

def get_text_from_article_url(url: str):
    """기사 URL에서 본문 텍스트를 추출합니다."""
    print(f"    (File Handler) - Extracting text from article: {url}")
    try:
        headers = {'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/58.0.3029.110 Safari/537.3'}
        response = requests.get(url, headers=headers, timeout=10)
        response.raise_for_status()
        soup = BeautifulSoup(response.text, 'html.parser')
        
        title = soup.find('h1').get_text(strip=True) if soup.find('h1') else 'No Title Found'
        article_body = soup.find('article') or soup.find('main') or soup.find('body')
        paragraphs = article_body.find_all('p') if article_body else []
        body_text = "\n".join([p.get_text() for p in paragraphs])
        
        full_content = f"Article Title: {title}\n\nBody:\n{body_text}"
        return full_content, "text"
    except Exception as e:
        print(f"    [ERROR] Failed to process article link: {e}")
        return None, None

def get_transcript_from_youtube_url(url: str):
    """유튜브 URL에서 스크립트(자막)를 텍스트로 추출합니다."""
    print(f"    (File Handler) - Extracting transcript from YouTube: {url}")
    try:
        if 'v=' in url:
            video_id = url.split('v=')[-1].split('&')[0]
        elif 'youtu.be/' in url:
            video_id = url.split('youtu.be/')[-1].split('?')[0]
        else:
            raise ValueError("Invalid YouTube URL")
            
        transcript_list = YouTubeTranscriptApi.get_transcript(video_id, languages=['ko', 'en'])
        transcript = " ".join([item['text'] for item in transcript_list])
        return transcript, "text"
    except Exception as e:
        print(f"    [ERROR] Failed to process YouTube link: {e}")
        return None, None

def get_content_from_input(path_or_url: str):
    """
    URL, 로컬 파일, 또는 로컬 디렉터리 경로를 읽어 콘텐츠와 타입을 반환합니다.
    """
    print(f"    (File Handler) - Processing input: {path_or_url}")

    # 1. URL 처리
    if path_or_url.startswith(('http://', 'https://')):
        return get_transcript_from_youtube_url(path_or_url) if ("youtube.com" in path_or_url or "youtu.be" in path_or_url) else get_text_from_article_url(path_or_url)

    file_path = path_or_url

    # 2. 로컬 디렉터리 경로 처리 (로컬 환경 전용)
    if os.path.isdir(file_path):
        print(f"    (File Handler) - Path is a directory. Searching for files.")
        try:
            files_in_dir = [f for f in os.listdir(file_path) if os.path.isfile(os.path.join(file_path, f))]
            if not files_in_dir:
                return None, None
            file_path = os.path.join(file_path, files_in_dir[0])
            print(f"    (File Handler) - Processing first file found: {file_path}")
        except Exception as e:
            print(f"    [ERROR] Failed to list files in directory {file_path}: {e}")
            return None, None

    # 3. 파일 경로 처리 (클라우드 및 로컬)
    if not os.path.exists(file_path):
        print(f"    [ERROR] File not found: {file_path}")
        return None, None

    try:
        file_extension = os.path.splitext(file_path)[1].lower()
        print(f"    (File Handler) - Detected file type: '{file_extension}'")

        # --- 이미지 파일은 경로를 그대로 반환 ---
        if file_extension in ['.jpg', '.jpeg', '.png']:
            return file_path, "file_path"

        # --- 텍스트 기반 파일들은 내용 추출 ---
        content = ""
        if file_extension == '.csv':
            df = pd.read_csv(file_path, encoding='utf-8', on_bad_lines='skip')
            content = df.to_string()
        elif file_extension == '.xlsx':
            df = pd.read_excel(file_path)
            content = df.to_string()
        elif file_extension == '.txt':
            with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
                content = f.read()
        elif file_extension == '.pdf':
            with open(file_path, 'rb') as f:
                reader = PyPDF2.PdfReader(f, strict=False)
                for page in reader.pages: content += page.extract_text() or ""
        elif file_extension == '.docx':
            doc = docx.Document(file_path)
            content = "\n".join([para.text for para in doc.paragraphs if para.text])
        elif file_extension == '.pptx':
            prs = Presentation(file_path)
            for slide in prs.slides:
                for shape in slide.shapes:
                    if hasattr(shape, "text"): content += shape.text + "\n"
        
        return content, "text"

    except Exception as e:
        print(f"    [ERROR] Failed to process file {file_path}: {e}")
        return None, None